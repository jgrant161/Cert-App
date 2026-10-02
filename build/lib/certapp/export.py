"""Excel deliverable: the certificate schedule plus the supporting tabs."""

from __future__ import annotations

import io
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .rules import POLICY_LABELS, EngagementResult, Policy
from .states import STATES

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
FILLS = {
    "tbd_resolved": PatternFill("solid", fgColor="D9EAD3"),   # green: resolved from document
    "corrected": PatternFill("solid", fgColor="FCE5CD"),      # orange: filename was wrong
    "override": PatternFill("solid", fgColor="CFE2F3"),       # blue: reviewer override
    "flagged": PatternFill("solid", fgColor="FFF2CC"),        # yellow: needs a look
}
BASIS_LABELS = {
    "single_state": "Single-state certificate",
    "exploded": "Multistate form (one row per state)",
    "tbd_resolved": "State resolved from document (filename TBD)",
    "corrected": "State corrected from document",
    "override": "Reviewer override",
}


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], widths: list[int] | None = None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for i, h in enumerate(headers, start=1):
        width = (widths[i - 1] if widths else None) or max(
            [len(str(h))] + [len(str(r[i - 1] or "")) for r in rows[:500]]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 60)
    return ws


def build_workbook(engagement: dict, result: EngagementResult, policy: Policy,
                   batches: list[dict]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    rows = sorted(result.rows, key=lambda r: (r.company.lower(), r.state))

    # Schedule
    ws = _sheet(wb, "Schedule",
                ["Company", "Certificate Type", "State", "State Name", "Registration / ID",
                 "Form Number", "Form Name", "Basis", "Batch", "Source File", "Flags", "Notes"],
                [[r.company, r.certificate_type, r.state, STATES.get(r.state, ""), r.registration,
                  r.form_number, r.form_name, BASIS_LABELS.get(r.basis, r.basis), r.batch, r.source_file,
                  ", ".join(dict.fromkeys(r.flags)), r.note] for r in rows],
                [36, 22, 7, 16, 22, 14, 40, 30, 16, 50, 34, 40])
    for i, r in enumerate(rows, start=2):
        fill = FILLS.get(r.basis) or (FILLS["flagged"] if r.flags else None)
        if fill:
            for cell in ws[i]:
                cell.fill = fill

    # Summaries
    by_state = result.by_state()
    _sheet(wb, "Summary by State", ["State", "State Name", "Rows"],
           [[s, STATES.get(s, ""), n] for s, n in by_state] + [["Total", "", sum(n for _, n in by_state)]])
    _sheet(wb, "Summary by Company", ["Company", "State Count", "States"],
           [list(t) for t in result.by_company()], [40, 12, 80])

    # Follow-up: everything that is not purely informational
    follow = []
    for c in result.certificates:
        for f in c.flags:
            if f.severity == "info" and f.code != "entity_based":
                continue
            follow.append([f.severity.replace("_", " ").title(), f.code.replace("_", " "),
                           c.cert["filename"], f.message, c.cert.get("batch_label") or ""])
    order = {"Follow Up": 0, "Review": 1, "Info": 2}
    follow.sort(key=lambda r: (order.get(r[0], 3), r[2].lower()))
    _sheet(wb, "Follow-Up", ["Severity", "Issue", "Source File", "Detail", "Batch"], follow,
           [12, 24, 50, 90, 16])

    # Lines present on a form but kept off the schedule
    _sheet(wb, "Excluded Lines", ["Company", "State", "Text on Line", "Classification", "Why Excluded", "Source File"],
           [[x.company, x.state, x.raw_text, x.classification.replace("_", " "), x.reason, x.source_file]
            for x in sorted(result.excluded, key=lambda x: (x.company.lower(), x.state))],
           [36, 7, 40, 20, 50, 50])

    # Reconciliation: every source file and what it produced
    recon = []
    for c in sorted(result.certificates, key=lambda c: c.cert["filename"].lower()):
        ex = c.extraction
        recon.append([c.cert["filename"], c.cert.get("batch_label") or "", c.cert["status"],
                      ex.form_number if ex else None, ex.form_name if ex else None,
                      len(c.rows), ", ".join(sorted({r.state for r in c.rows})),
                      c.cert.get("review_status") or "", c.cert.get("reviewer_note") or ""])
    total_rows = sum(r[5] for r in recon)
    recon.append([f"TOTAL ({len(recon)} files)", "", "", "", "", total_rows, "", "", ""])
    _sheet(wb, "Reconciliation", ["Source File", "Batch", "Status", "Form Number", "Form Name", "Rows",
                                  "States", "Review", "Reviewer Note"],
           recon, [60, 16, 12, 14, 40, 8, 50, 12, 50])

    statuses = Counter(c.cert["status"] for c in result.certificates)
    _sheet(wb, "Batches", ["Batch", "Uploaded", "Files Received", "Added", "Already On File"],
           [[b["label"], b["uploaded_at"], b["files_received"], b["files_new"], b["files_duplicate"]]
            for b in batches])

    overview = [
        ["Client", engagement["client_name"]],
        ["Seller", engagement.get("seller_name") or ""],
        ["Certificates on file", len(result.certificates)],
        ["Read", statuses.get("extracted", 0)],
        ["Not yet read", statuses.get("pending", 0) + statuses.get("processing", 0)],
        ["Errors", statuses.get("error", 0)],
        ["Schedule rows", len(rows)],
        ["States", len(by_state)],
        ["", ""],
        ["Review policy", ""],
    ] + [[label, "Yes" if getattr(policy, key) else "No"] for key, label in POLICY_LABELS.items()]
    ov = _sheet(wb, "Overview", ["Item", "Value"], overview, [70, 40])
    wb.move_sheet(ov, offset=-(len(wb.sheetnames) - 1))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
