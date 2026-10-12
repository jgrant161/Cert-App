"""End to end through the web app: upload a zip, read, review, add a batch, export."""

import io
import zipfile

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from certapp.extraction.schema import json_schema
from certapp.web.app import create_app

from conftest import FakeExtractor, admin_client, extraction, line, make_pdf

FILES = {
    "Acme Supply - Resale Certificate (CA).pdf": extraction(
        [line("CA", "SR-111")], purchaser="Acme", form_number="CDTFA-230", form="California Resale Certificate"),
    "Brandster - Resale Certificate (Multijurisdiction).pdf": extraction(
        [line("FL", "FL-7"), line("AL", "FL-7", "home_state_number", issuing="FL"),
         line("GA", "", "blank")], multistate=True, purchaser="Brandster"),
    "Collins - Resale Certificate (FL).pdf": extraction([line("NC", "600123")], purchaser="Collins"),
}
LATER = {"Delta Tools - Resale Certificate (TX).pdf": extraction([line("TX", "3200")], purchaser="Delta")}


def zip_of(names):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for n in names:
            zf.writestr(f"Certificates/{n}", make_pdf(n))
        zf.writestr("__MACOSX/._junk", b"x")
    return buf.getvalue()


def test_full_engagement_flow():
    fake = FakeExtractor({**FILES, **LATER})
    answers = []
    client = admin_client(TestClient(create_app(extractor=fake, ask_fn=lambda q, ctx, h: answers.append(ctx) or "answer")))

    r = client.post("/engagements", data={"client_name": "Additive Manufacturing, LLC"}, follow_redirects=False)
    eid = int(r.headers["location"].rsplit("/", 1)[1])

    r = client.post(f"/e/{eid}/upload", data={"label": "Initial"},
                    files=[("files", ("Certificates.zip", zip_of(FILES), "application/zip"))])
    assert r.status_code == 200
    assert sorted(fake.calls) == sorted(FILES)

    page = client.get(f"/e/{eid}").text
    assert "SR-111" in page and "AL" in page and "CDTFA-230" in page
    assert "California Resale Certificate" in client.get(f"/e/{eid}?tab=certificates").text

    # Second batch: one new file, one re-sent file (skipped, not re-read).
    resend = zip_of(["Acme Supply - Resale Certificate (CA).pdf", *LATER])
    client.post(f"/e/{eid}/upload", data={"label": "9.30.26 Certs"},
                files=[("files", ("9.30.26 Certs.zip", resend, "application/zip"))])
    assert fake.calls.count("Acme Supply - Resale Certificate (CA).pdf") == 1
    assert "Delta Tools - Resale Certificate (TX).pdf" in fake.calls
    batches = client.get(f"/e/{eid}?tab=batches").text
    assert "9.30.26 Certs" in batches

    # Reviewer excludes Brandster's AL line.
    cert_page = client.get(f"/e/{eid}/c/2")
    assert cert_page.status_code == 200 and "Home-state number" in cert_page.text
    client.post(f"/e/{eid}/c/2/override", data={"state": "AL", "decision": "exclude", "note": "per client"})
    client.post(f"/e/{eid}/c/2/review", data={"review_status": "reviewed", "note": "checked"})

    # Policy page renders and saves.
    assert client.get(f"/e/{eid}?tab=policy").status_code == 200
    client.post(f"/e/{eid}/policy", data={"include_fein": "on", "expiring_soon_days": "60"})

    # Follow-up tab shows the wrong-state filename.
    assert "filename state mismatch" in client.get(f"/e/{eid}?tab=followup").text

    # Excel deliverable
    r = client.get(f"/e/{eid}/export.xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[0] == "Overview"
    sched = [row for row in wb["Schedule"].iter_rows(min_row=2, values_only=True)]
    header = [c.value for c in wb["Schedule"][1]]
    acme = dict(zip(header, next(row for row in sched if row[0] == "Acme Supply")))
    assert (acme["Form Number"], acme["Form Name"]) == ("CDTFA-230", "California Resale Certificate")
    assert sorted((row[0], row[3]) for row in sched) == [
        ("Acme Supply", "CA"), ("Brandster", "FL"), ("Collins", "NC"), ("Delta Tools", "TX")]
    assert {"Status", "Expires", "Date Signed", "Seller Match", "Reviewer Note", "Flags"} <= set(header)
    rh = [c.value for c in wb["Reconciliation"][1]]
    recon = [dict(zip(rh, row)) for row in wb["Reconciliation"].iter_rows(min_row=2, values_only=True)]
    assert recon[-1]["Source File"] == "TOTAL (4 certificates)" and recon[-1]["Rows"] == 4
    acme_recon = next(r for r in recon if r["Source File"].startswith("Acme"))
    assert (acme_recon["Form Number"], acme_recon["Form Name"]) == ("CDTFA-230", "California Resale Certificate")
    assert any(row[0] == "Brandster" and row[1] == "AL" for row in
               wb["Excluded Lines"].iter_rows(min_row=2, values_only=True))

    # Q&A gets the engagement data as context (demo mode answers without calling out).
    client.post(f"/e/{eid}/ask", data={"question": "Which files were renamed?"})
    assert "Which files were renamed?" in client.get(f"/e/{eid}?tab=ask").text


def test_extraction_errors_are_isolated_and_retryable():
    fake = FakeExtractor({})
    client = admin_client(TestClient(create_app(extractor=fake)))
    client.post("/engagements", data={"client_name": "X"})
    client.post("/e/1/upload", files=[("files", ("Bad - Resale Certificate (CA).pdf", make_pdf("bad"), "application/pdf"))])
    assert "unreadable" in client.get("/e/1/c/1").text
    fake.by_filename["Bad - Resale Certificate (CA).pdf"] = extraction([line("CA", "SR-1")])
    client.post("/e/1/c/1/retry")
    assert "SR-1" in client.get("/e/1").text


def test_readings_saved_before_form_number_still_load():
    from certapp.extraction.schema import Extraction
    old = extraction([line("CA", "SR-1")]).model_dump()
    del old["form_number"]
    assert Extraction.model_validate(old).form_number is None


def test_json_schema_is_strict():
    schema = json_schema()

    def walk(node):
        if isinstance(node, dict):
            if node.get("type") == "object" and "properties" in node:
                assert node["additionalProperties"] is False
                assert set(node["required"]) == set(node["properties"])
            for v in node.values():
                walk(v)
        elif isinstance(node, list):
            for v in node:
                walk(v)
    walk(schema)


def test_upload_summary_and_delete_engagement(data_dir):
    fake = FakeExtractor(FILES)
    client = admin_client(TestClient(create_app(extractor=fake)))
    client.post("/engagements", data={"client_name": "Acme Client"})
    names = list(FILES)
    zip_bytes = zip_of(names)
    r = client.post("/e/1/upload", data={"label": "First"},
                    files=[("files", ("c.zip", zip_bytes, "application/zip")),
                           ("files", ("notes.txt", b"hello", "text/plain"))])
    assert 'Batch "First" uploaded.' in r.text and "3 new certificates added" in r.text
    assert "1 non-PDF file ignored" in r.text
    r = client.post("/e/1/upload", data={"label": "Again"},
                    files=[("files", ("c.zip", zip_bytes, "application/zip"))])
    assert "0 new certificates added" in r.text and "3 already on file (skipped)" in r.text

    # Wrong name: nothing deleted.
    r = client.post("/e/1/delete", data={"confirm_name": "Acme"})
    assert "did not match" in r.text
    assert client.get("/e/1").status_code == 200

    files_dir = data_dir / "data" / "files" / "1"
    assert any(files_dir.iterdir())
    r = client.post("/e/1/delete", data={"confirm_name": " acme client "})
    assert 'Deleted "Acme Client"' in r.text
    assert client.get("/e/1").status_code == 404
    assert not files_dir.exists()


def test_interrupted_reads_resume():
    from certapp import db
    eid = db.create_engagement("X")
    name = "Acme Supply - Resale Certificate (CA).pdf"
    from certapp.ingest import ingest
    ingest(eid, [(name, make_pdf("x"))], "b")
    db.claim_pending(eid)                       # app "closed" mid-read
    fake = FakeExtractor(FILES)
    client = admin_client(TestClient(create_app(extractor=fake)))   # restart resets it
    page = client.get(f"/e/{eid}").text
    assert "1 certificate waiting to be read" in page and "Resume reading" in page
    client.post(f"/e/{eid}/process")
    assert "SR-111" in client.get(f"/e/{eid}").text


def test_bundled_file_is_split_into_certificates():
    """One PDF with two state certificates and a W-9 becomes two rows, each tied to its pages."""
    from certapp import db
    from certapp.extraction.schema import OtherPage, Reading
    bundle = "Renovation Brands - Resale Certificates (Multiple).pdf"
    reading = Reading(certificates=[
        extraction([line("MA", "ST4-1")], purchaser="Renovation Brands").model_copy(update={"source_pages": [1]}),
        extraction([line("CT", "CT-22")], purchaser="Renovation Brands").model_copy(update={"source_pages": [3, 4]}),
    ], other_pages=[OtherPage(pages=[2], kind="w9", description="IRS Form W-9")])
    fake = FakeExtractor({bundle: reading})
    client = admin_client(TestClient(create_app(extractor=fake)))
    eid = db.create_engagement("Bundles")
    client.post(f"/e/{eid}/upload", files=[("files", (bundle, make_pdf("bundle"), "application/pdf"))])

    certs = db.list_certificates(eid)
    assert [(c["filename"], c["pages"]) for c in certs] == [
        (bundle, "1"), (f"{bundle} (pages 3-4)", "3-4")]
    child = certs[1]
    assert child["parent_id"] == certs[0]["id"] and child["sha256"] == certs[0]["sha256"]
    page = client.get(f"/e/{eid}/c/{certs[0]['id']}").text
    assert "IRS Form W-9" in page and "held 2 certificates" in page
    assert "Split out of" in client.get(f"/e/{eid}/c/{child['id']}").text
    sched = client.get(f"/e/{eid}").text
    assert "ST4-1" in sched and "CT-22" in sched

    # Sending the same file again adds nothing; re-reading the part re-reads the whole file.
    client.post(f"/e/{eid}/upload", files=[("files", (bundle, make_pdf("bundle"), "application/pdf"))])
    assert len(db.list_certificates(eid)) == 2
    fake.by_filename[bundle] = reading.certificates[0]
    r = client.post(f"/e/{eid}/c/{child['id']}/retry", follow_redirects=False)
    assert r.headers["location"] == f"/e/{eid}/c/{certs[0]['id']}"
    assert [c["filename"] for c in db.list_certificates(eid)] == [bundle]


def test_file_with_no_certificate_is_an_error():
    from certapp import db
    from certapp.extraction.schema import OtherPage, Reading
    name = "Cover - Resale Certificate (CA).pdf"
    fake = FakeExtractor({name: Reading(certificates=[], other_pages=[
        OtherPage(pages=[1], kind="cover_letter", description="Cover letter from the customer")])})
    client = admin_client(TestClient(create_app(extractor=fake)))
    eid = db.create_engagement("Cover")
    client.post(f"/e/{eid}/upload", files=[("files", (name, make_pdf("c"), "application/pdf"))])
    c = db.list_certificates(eid)[0]
    assert c["status"] == "error" and "Cover letter" in c["error"]


def test_review_status_override_policy_settings_and_customer_merge():
    from certapp import db
    fake = FakeExtractor({**FILES, "Acme Supply Inc - Resale Certificate (TX).pdf": extraction(
        [line("TX", "3200")], purchaser="ACME Supply, Inc.")})
    client = admin_client(TestClient(create_app(extractor=fake)))
    eid = db.create_engagement("Merge Co")
    client.post(f"/e/{eid}/upload", files=[("files", (n, make_pdf(n), "application/pdf")) for n in fake.by_filename])
    page = client.get(f"/e/{eid}?tab=customers").text
    assert page.count("<b>Acme Supply") == 1                         # name variants joined automatically

    client.post(f"/e/{eid}/customers/merge", data={"source": "Collins", "target": "Brandster"})
    assert "<b>Collins</b>" not in client.get(f"/e/{eid}?tab=customers").text
    client.post(f"/e/{eid}/policy", data={"include_fein": "on", "expiration_rules": "IL: 3 years",
                                           "uniform_not_accepted": "HI, MD"})
    pol = client.get(f"/e/{eid}?tab=policy").text
    assert "3 years from the signature date" in pol and 'value="HI, MD"' in pol
    from certapp.rules import Policy
    assert Policy.from_json(db.get_engagement(eid)["policy_json"]).customer_aliases == {"collins": "brandster"}
    client.post(f"/e/{eid}/customers/merge", data={"source": "Collins", "target": ""})
    assert "<b>Collins</b>" in client.get(f"/e/{eid}?tab=customers").text

    cid = db.list_certificates(eid)[0]["id"]
    client.post(f"/e/{eid}/c/{cid}/review", data={"review_status": "reviewed", "status_override": "invalid",
                                                   "note": "Wrong entity"})
    assert db.get_certificate(cid)["status_override"] == "invalid"
    assert "status set by a reviewer" in client.get(f"/e/{eid}/c/{cid}").text
    assert client.post(f"/e/{eid}/c/{cid}/review", data={"review_status": "bogus"}).status_code == 400


def test_older_database_gains_new_columns(data_dir):
    import os
    import sqlite3
    import pytest
    from certapp import db
    if os.environ.get("DATABASE_URL"):
        pytest.skip("SQLite upgrade path")
    path = db.data_dir() / "certapp.db"
    path.unlink()
    with sqlite3.connect(path) as raw:              # the certificate table as first released
        raw.execute("""CREATE TABLE certificate (id INTEGER PRIMARY KEY, engagement_id INTEGER NOT NULL,
            batch_id INTEGER NOT NULL, filename TEXT NOT NULL, sha256 TEXT NOT NULL, storage_path TEXT NOT NULL,
            filename_json TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'pending', extractor TEXT,
            extraction_json TEXT, error TEXT, duplicate_of INTEGER, review_status TEXT NOT NULL DEFAULT 'open',
            reviewer_note TEXT, created_at TEXT NOT NULL)""")
    db.init()
    with sqlite3.connect(path) as raw:
        cols = {r[1] for r in raw.execute("PRAGMA table_info(certificate)")}
    assert {"status_override", "other_pages_json", "pages", "parent_id"} <= cols
